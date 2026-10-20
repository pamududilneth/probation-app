"""
Fills the ORIGINAL Word template (word/document.xml, untouched layout) with
the current form data and returns the resulting .docx as bytes.

Approach: rather than regenerating the document from scratch, we open the
exact uploaded template and insert new runs/paragraphs/images at precise,
manually mapped locations (row/cell/paragraph indices verified against the
real file) — so every label, merged cell, and column width stays exactly as
authored. Signature fields arrive as actual drawn images (canvas PNG data
URLs from the frontend), not typed names, and are embedded as real inline
pictures. The HOD/CEO approval boxes are Word floating textboxes — we edit
their "Signature:" / "Date:" / decision-checkbox text directly, in place,
rather than writing a summary elsewhere.
"""
import base64
import io
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import Pt

TEMPLATE_PATH = Path(__file__).parent / "templates" / "probation_template.docx"

FACTORS_ORDER = [
    "job_knowledge", "quality_of_work", "responsibility", "dependability",
    "courtesy", "personal_tidiness", "relationships", "attendance",
    "initiative", "company_property", "overall_eval",
]

INTERVAL_LABELS = {"6": "6 Months", "12": "12 Months", "end": "Probationary Period End"}
RESULT_LABELS = {
    "contributor": "Contributor",
    "extended": "Probationary Period Extended",
    "not-suitable": "Not Suitable for Employment",
}
DECISION_LABELS = {"confirm": "Confirm", "extend": "Extend Probation", "not-suitable": "Not Suitable for Employment"}

GRADE_COLUMN = {"A": 2, "B": 4, "C": 5, "D": 6}
COMMENT_COLUMN = 7


def _insert_after(paragraph, match_substr, value, bold=False, skip=0):
    if not value:
        return False
    runs = paragraph.runs
    for i, run in enumerate(runs):
        if match_substr in run.text:
            anchor = runs[min(i + skip, len(runs) - 1)]
            new_run = paragraph.add_run(f" {value}")
            new_run.bold = bold
            anchor._element.addnext(new_run._element)
            return True
    return False


def _fill_blank_paragraph(cell, index, value):
    if not value:
        return
    if index < len(cell.paragraphs):
        cell.paragraphs[index].add_run(str(value))
    else:
        cell.add_paragraph(str(value))


def _decode_signature(value):
    """Signature fields arrive from the frontend as a canvas-drawn PNG data
    URL (e.g. 'data:image/png;base64,....'), not typed text — a typed name
    isn't a signature. Returns raw image bytes, or None if there's nothing
    (or something malformed) to embed."""
    if not value or not isinstance(value, str) or not value.startswith("data:image"):
        return None
    try:
        _, b64data = value.split(",", 1)
        return base64.b64decode(b64data)
    except Exception:
        return None


def _insert_image_after_label(paragraph, label_marker, image_bytes, height_pt=26):
    """Splits the run containing `label_marker` (e.g. 'Signature:') right
    after the label, inserts the actual signature image there, and puts
    whatever text originally followed the label back afterwards unchanged."""
    if not image_bytes:
        return False
    for run in list(paragraph.runs):
        if label_marker in run.text:
            idx = run.text.find(label_marker) + len(label_marker)
            before, after = run.text[:idx], run.text[idx:]
            run.text = before
            img_run = paragraph.add_run()
            img_run.add_picture(io.BytesIO(image_bytes), height=Pt(height_pt))
            run._element.addnext(img_run._element)
            if after:
                after_run = paragraph.add_run(after)
                after_run.bold = run.bold
                img_run._element.addnext(after_run._element)
            return True
    return False


def _find_drawing_signature_paragraphs(doc, container_elem, marker="Signature:"):
    """The HOD/CEO approval boxes are Word floating textboxes, which Word
    duplicates internally (once for modern rendering, once for legacy
    compatibility) — so the same visible line can exist as 2-4 near-
    identical paragraph copies in the XML. We edit every copy we find so
    the result is correct regardless of which one a given Word version
    actually renders."""
    from docx.text.paragraph import Paragraph

    out = []
    for txbx in container_elem.findall(".//" + qn("w:txbxContent")):
        for p in txbx.findall(qn("w:p")):
            if marker in "".join(p.itertext()):
                out.append(Paragraph(p, doc))
    return out


def _mark_decision(paragraphs, decision, extend_duration=""):
    """Ticks off the chosen option (Confirm / Extend Probation by... /
    Not Suitable for Employment) directly inside the approval box."""
    if not decision:
        return
    if decision == "confirm":
        for p in paragraphs:
            _insert_after(p, "Confirm", "\u2713", bold=True)
    elif decision == "extend":
        for p in paragraphs:
            _insert_after(p, "Extend Probation by", "\u2713", bold=True)
    elif decision == "not-suitable":
        for p in paragraphs:
            _insert_after(p, "Not Suitable for Employment", "\u2713", bold=True)


def _replace_placeholder_runs(container_paragraphs, full_name, epf):
    for p in container_paragraphs:
        for run in p.runs:
            stripped = run.text.strip()
            if stripped == "Full Name":
                run.text = full_name or "________"
            elif stripped in ("EPF Num", "EFP Num"):
                run.text = epf or "________"


def _walk_all_paragraphs(doc):
    yield from doc.paragraphs

    def walk_tables(tables):
        for t in tables:
            for row in t.rows:
                for cell in row.cells:
                    yield from cell.paragraphs
                    if cell.tables:
                        yield from walk_tables(cell.tables)

    yield from walk_tables(doc.tables)


def generate_docx(data: dict) -> bytes:
    doc = docx.Document(str(TEMPLATE_PATH))

    full_name = data.get("fullName", "")
    epf = data.get("epfTop") or data.get("epfNum", "")

    _replace_placeholder_runs(list(_walk_all_paragraphs(doc)), full_name, epf)

    t = doc.tables[0]

    # Row 1 — Projected Probationary End Date: []
    for run in t.rows[1].cells[0].paragraphs[0].runs:
        if run.text.strip() == "[]":
            run.text = f"[{data.get('probEndDate', '')}]"

    # Row 2 — Review Interval
    interval = data.get("interval")
    if interval:
        label = INTERVAL_LABELS.get(interval) or data.get("intervalOther", "")
        note = t.rows[2].cells[0].add_paragraph()
        run = note.add_run(f"Selected: {label}")
        run.bold = True
        run.italic = True

    # Row 3 — EPF Number
    _insert_after(t.rows[3].cells[0].paragraphs[0], "EPF Number:", data.get("epfNum", ""))

    # Row 4 — Department / Designation
    p4 = t.rows[4].cells[0].paragraphs[0]
    _insert_after(p4, "Department:", data.get("dept", ""))
    _insert_after(p4, "Designation:", data.get("designation", ""))

    # Row 5 — Employment Date
    _insert_after(t.rows[5].cells[0].paragraphs[0], "Employment Date", data.get("empDate", ""), skip=1)

    # Row 6 — Supervisor's Name / Title
    p6 = t.rows[6].cells[0].paragraphs[0]
    _insert_after(p6, "Supervisor\u2019s Name", data.get("supName", ""), skip=1)
    _insert_after(p6, "Supervisor\u2019s Title", data.get("supTitle", ""), skip=1)

    # Row 7 — Leave records
    p7_1 = t.rows[7].cells[0].paragraphs[1]
    _insert_after(p7_1, "Absent without leave", data.get("leaveAwol", ""))
    _insert_after(p7_1, "No Pay", data.get("leaveNoPay", ""))
    p7_2 = t.rows[7].cells[0].paragraphs[2]
    _insert_after(p7_2, "Annual", data.get("leaveAnnual", ""))
    _insert_after(p7_2, "Casual", data.get("leaveCasual", ""))
    _insert_after(p7_2, "Medical", data.get("leaveMedical", ""))

    # Row 9 — Employee Development Plan
    _fill_blank_paragraph(t.rows[9].cells[0], 1, data.get("devPlan", ""))
    # Row 10 — Steps taken to enhance job performance
    _fill_blank_paragraph(t.rows[10].cells[0], 1, data.get("devSteps", ""))

    # Row 11 — Employee signature / date
    p11 = t.rows[11].cells[0].paragraphs[0]
    _insert_image_after_label(p11, "Signature:", _decode_signature(data.get("empSig")))
    _insert_after(p11, "Date", data.get("empSigDate", ""))

    # Nested evaluation table
    nested = t.rows[13].cells[0].tables[0]
    for i, factor in enumerate(FACTORS_ORDER):
        row = nested.rows[i + 2]
        grade = data.get(f"grade_{factor}")
        comment = data.get(f"comment_{factor}", "")
        if grade in GRADE_COLUMN:
            cell = row.cells[GRADE_COLUMN[grade]]
            para = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
            run = para.add_run("X")
            run.bold = True
        if comment:
            cell = row.cells[COMMENT_COLUMN]
            para = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
            para.add_run(str(comment))

    # Row 13 cell — Comments / strengths / improvements
    cell13 = t.rows[13].cells[0]
    _fill_blank_paragraph(cell13, 5, data.get("overallProgress", ""))
    _fill_blank_paragraph(cell13, 9, data.get("strengths", ""))
    _fill_blank_paragraph(cell13, 13, data.get("improvements", ""))

    # Row 14 — Overall Result of Review
    result = data.get("overallResult")
    if result:
        label = RESULT_LABELS.get(result, result)
        extra = f" (until {data.get('extendedUntil', '')})" if result == "extended" else ""
        note = t.rows[14].cells[0].add_paragraph()
        run = note.add_run(f"Selected: {label}{extra}")
        run.bold = True
        run.italic = True

    # Row 16 — Steps taken to assist employee
    _fill_blank_paragraph(t.rows[16].cells[0], 1, data.get("assistSteps", ""))

    # Row 17 — Supervisor signature / date
    p17 = t.rows[17].cells[0].paragraphs[0]
    _insert_image_after_label(p17, "Signature:", _decode_signature(data.get("supSig")))
    _insert_after(p17, "Date", data.get("supSigDate", ""))

    # HOD / CEO decisions — edited directly inside the approval boxes
    # themselves, not written as a separate summary elsewhere.
    hod, ceo = data.get("hodDecision"), data.get("ceoDecision")
    body_children = list(doc.element.body)
    tbl_elems = [c for c in body_children if c.tag == qn("w:tbl")]
    idx_table0 = body_children.index(tbl_elems[0])

    hod_container = body_children[idx_table0 + 1]
    ceo_container = None
    for c in body_children[idx_table0 + 1:]:
        if c.tag == qn("w:p") and "CEO" in "".join(c.itertext()):
            ceo_container = c
            break

    for container, decision, extend_key, sig_key, sig_date_key in [
        (hod_container, hod, "hodExtendDuration", "hodSig", "hodSigDate"),
        (ceo_container, ceo, "ceoExtendDuration", "ceoSig", "ceoSigDate"),
    ]:
        if container is None:
            continue
        sig_paragraphs = _find_drawing_signature_paragraphs(doc, container, marker="Signature:")
        decision_paragraphs = _find_drawing_signature_paragraphs(doc, container, marker="Confirm")
        if not sig_paragraphs:
            continue

        _mark_decision(decision_paragraphs, decision, data.get(extend_key, ""))

        sig_bytes = _decode_signature(data.get(sig_key))
        for p in sig_paragraphs:
            _insert_image_after_label(p, "Signature:", sig_bytes, height_pt=20)
            _insert_after(p, "Date", data.get(sig_date_key, ""))

    # Table 1 — SBU HR
    t1 = doc.tables[1]
    if data.get("sbuComments"):
        t1.rows[0].cells[0].add_paragraph(str(data["sbuComments"]))
    p_sbu = t1.rows[1].cells[0].paragraphs[0]
    _insert_image_after_label(p_sbu, "Signature:", _decode_signature(data.get("sbuSig")))
    _insert_after(p_sbu, "Date", data.get("sbuSigDate", ""))

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()